import asyncio
import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from app.services.pipeline.report_service import ReportService
from app.services.pipeline.schemas import (
    PatentResearchReport, ReportPatent, ReportPatentDetails, ReportPatentMethodology,
    ReportComparisonDimension, DynamicTargetAttributeEvidence, MediumAndWaterRoleEvidence,
    ExtractedTableSchema
)
from tests.test_pdf_report_layout import REFERENCE_DOCX, _report_from_reference_docx

async def main():
    service = ReportService()
    service.export_dir = "exports"
    os.makedirs(service.export_dir, exist_ok=True)
    
    print("=== 1. REGRESSION TEST: LOW STYRENE SBR (10 PATENTS) ===")
    if REFERENCE_DOCX.exists():
        sbr_report = _report_from_reference_docx(REFERENCE_DOCX)
        print(f"Methodology patents count: {len(sbr_report.methodology_patents)}")
        sbr_md = service.report_to_markdown(sbr_report)
        sbr_pdf = await service.export_to_pdf(sbr_md, "regression_sbr_report.pdf")
        sbr_docx = await service.export_to_docx(sbr_report, "regression_sbr_report.docx")
        
        pdf_reader = PdfReader(sbr_pdf)
        pdf_pages = len(pdf_reader.pages)
        print(f"SBR PDF Page Count: {pdf_pages}")
        
        docx_doc = Document(sbr_docx)
        docx_headings = [p.text for p in docx_doc.paragraphs if p.style.name.startswith("Heading")]
        print(f"SBR DOCX Headings: {docx_headings}")
        print(f"SBR DOCX Tables count: {len(docx_doc.tables)}")
        
        for i, page in enumerate(pdf_reader.pages):
            txt = page.extract_text() or ""
            lines = [l.strip() for l in txt.split("\n") if l.strip()]
            first_line = lines[0] if lines else "EMPTY"
            print(f"  Page {i+1} ({len(txt)} chars): starts with: {first_line[:60]}")
    else:
        print("REFERENCE_DOCX not found!")
        
    print("\n=== 2. TEST CASE: BIO-BASED PHA POLYMER (4 PATENTS) ===")
    pha_patents = [
        ReportPatent(
            patent_details=ReportPatentDetails(
                patent_number=f"WO2023{100000+i}A1",
                patent_title=f"Microbial synthesis of polyhydroxyalkanoate copolymer variant {i}",
                assignee="BioPolymers International Ltd.",
                jurisdiction="WO",
                publication_year="2023",
                relevance_to_target="Primary evidence on elongation at break"
            ),
            polymerization_method=ReportPatentMethodology(
                dynamic_parameters=[
                    f"Fermentation Temperature: {30 + i} °C",
                    "Carbon Substrate: Glucose / Sodium Valerate",
                    f"Cell Dry Weight: {80 + i*5} g/L",
                    f"Elongation at Break: {350 + i*50} %"
                ]
            ),
            experimental_evidence=[
                f"Example {i}: Fed-batch fermentation yielding high molecular weight PHA with elongation > 300%."
            ],
            technical_relevance="Demonstrates high elongation bio-polyester without synthetic plasticizers.",
            tables=[
                ExtractedTableSchema(
                    title=f"Table {i}: Fermentation Parameters",
                    headers=["Time (h)", "Biomass (g/L)", "PHA Accumulation (%)"],
                    rows=[["24", "45", "60"], ["48", "85", "78"]]
                )
            ],
            target_attribute=DynamicTargetAttributeEvidence(label="elongation at break", value=f"{350 + i*50} %")
        )
        for i in range(1, 5)
    ]
    pha_dims = [
        ReportComparisonDimension(
            parameter_name="elongation at break",
            values={p.patent_details.patent_number: f"{350 + i*50} %" for i, p in enumerate(pha_patents, 1)}
        ),
        ReportComparisonDimension(
            parameter_name="Cell Dry Weight",
            values={p.patent_details.patent_number: f"{80 + i*5} g/L" for i, p in enumerate(pha_patents, 1)}
        ),
    ]
    pha_report = PatentResearchReport(
        title="BIO-BASED PHA COPOLYMER PATENT RESEARCH REPORT",
        abstract="This report synthesizes technical intelligence across 4 selected patents regarding the microbial production and elongation properties of polyhydroxyalkanoate (PHA) copolymers. Selected disclosures reveal consistent optimization of valerate co-feed to enhance elongation at break above 300% while maintaining high cell biomass accumulation in fed-batch culture systems.",
        methodology_patents=pha_patents,
        cross_patent_comparison=[
            "Increasing sodium valerate co-feed ratio directly correlates with improved elongation at break across all selected patents.",
            "Fed-batch feeding schedules achieve over 75 wt% polymer accumulation inside bacterial cells."
        ],
        conclusion="The patent landscape reveals that microbial copolymerization with precise precursor dosing provides the dominant route to high-elongation PHA. Industrial scaling challenges center on carbon yield and purification efficiency without solvent intensive extraction.",
        references=[
            f"{p.patent_details.patent_number} | {p.patent_details.patent_title} | {p.patent_details.assignee} | {p.patent_details.jurisdiction} | 2023 | https://patents.google.com"
            for p in pha_patents
        ],
        comparison_dimensions=pha_dims
    )
    pha_md = service.report_to_markdown(pha_report)
    pha_pdf = await service.export_to_pdf(pha_md, "pha_copolymer_report.pdf")
    pha_docx = await service.export_to_docx(pha_report, "pha_copolymer_report.docx")
    
    pha_reader = PdfReader(pha_pdf)
    print(f"PHA PDF Page Count: {len(pha_reader.pages)}")
    pha_docx_doc = Document(pha_docx)
    heading_list = [p.text for p in pha_docx_doc.paragraphs if p.style.name.startswith("Heading")]
    print(f"PHA DOCX Headings: {heading_list}")
    print(f"PHA DOCX Tables count: {len(pha_docx_doc.tables)}")
    for i, page in enumerate(pha_reader.pages):
        txt = page.extract_text() or ""
        lines = [l.strip() for l in txt.split("\n") if l.strip()]
        first_line = lines[0] if lines else "EMPTY"
        print(f"  Page {i+1} ({len(txt)} chars): starts with: {first_line[:60]}")

if __name__ == "__main__":
    asyncio.run(main())
